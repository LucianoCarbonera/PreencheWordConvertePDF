import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from datetime import datetime
import locale
import os
import re
from docx2pdf import convert

# Definir localidade para português brasileiro
locale.setlocale(locale.LC_TIME, 'pt_BR.UTF-8')


def find_zip_files(directory):
    """Encontra todos os arquivos ZIP no diretório especificado."""
    return [os.path.join(directory, file_name) for file_name in os.listdir(directory) if file_name.endswith('.zip')]


def extract_xml_from_zip(zip_path, extract_to):
    """Extrai todos os arquivos XML de um ZIP e retorna os caminhos dos XMLs extraídos."""
    xml_files = []
    with zipfile.ZipFile(zip_path, 'r') as zip_ref:
        for file in zip_ref.namelist():
            if file.endswith('.xml'):
                zip_ref.extract(file, extract_to)
                xml_files.append(os.path.join(extract_to, file))
    return xml_files


def parse_xml_data(xml_path):
    """Parsa os dados do XML e retorna um dicionário com os dados relevantes."""
    try:
        tree = ET.parse(xml_path)
        root = tree.getroot()

        # Extraindo namespace
        ns = {'ns': root.tag.split('}')[0].strip('{')}

        # Extraindo dados do destinatário
        dest = root.find('.//ns:dest', ns)
        enderDest = dest.find('ns:enderDest', ns)
        nfe = root.find('.//ns:ide', ns)

        if dest is None or enderDest is None or nfe is None:
            raise ValueError("Estrutura do XML não corresponde ao esperado.")

        cpf = dest.find('ns:CPF', ns)
        cnpj = dest.find('ns:CNPJ', ns)
        if cpf is not None:
            identificacao = f"CPF: {format_cpf(cpf.text)}"
        elif cnpj is not None:
            identificacao = f"CNPJ: {format_cnpj(cnpj.text)}"
        else:
            identificacao = "N/A"

        data = {
            'Nome': dest.find('ns:xNome', ns).text if dest.find('ns:xNome', ns) is not None else "N/A",
            'Identificacao': identificacao,
            'Endereco': f"{enderDest.find('ns:xLgr', ns).text if enderDest.find('ns:xLgr', ns) is not None else ''}, {enderDest.find('ns:nro', ns).text if enderDest.find('ns:nro', ns) is not None else ''} - {enderDest.find('ns:xCpl', ns).text if enderDest.find('ns:xCpl', ns) is not None else ''}, {enderDest.find('ns:xBairro', ns).text if enderDest.find('ns:xBairro', ns) is not None else ''}, {enderDest.find('ns:xMun', ns).text if enderDest.find('ns:xMun', ns) is not None else ''} - {enderDest.find('ns:UF', ns).text if enderDest.find('ns:UF', ns) is not None else ''}",
            'CEP': f"{enderDest.find('ns:CEP', ns).text[:5]}-{enderDest.find('ns:CEP', ns).text[5:]}" if enderDest.find('ns:CEP', ns) is not None else "N/A",
            'NumeroNotaFiscal': nfe.find('ns:nNF', ns).text if nfe.find('ns:nNF', ns) is not None else "N/A"
        }

        # Extraindo dados do produto e número de série
        prod = root.find('.//ns:prod', ns)
        vol = root.find('.//ns:vol', ns)
        if prod is None or vol is None:
            raise ValueError("Dados do produto ou número de série não encontrados no XML.")

        data['ModeloGPS'] = prod.find('ns:xProd', ns).text if prod.find('ns:xProd', ns) is not None else "N/A"
        data['NumeroSerie'] = vol.find('ns:nVol', ns).text if vol.find('ns:nVol', ns) is not None else "N/A"

        return data
    except ET.ParseError:
        raise ValueError("Erro ao analisar o XML.")
    except Exception as e:
        raise ValueError(f"Erro ao extrair dados do XML: {e}")


def replace_placeholders(doc_path, output_path, replacements):
    """Substitui os placeholders no documento Word com os dados fornecidos."""
    try:
        doc = Document(doc_path)
        for paragraph in doc.paragraphs:
            for key, value in replacements.items():
                if key in paragraph.text:
                    paragraph.text = paragraph.text.replace(key, value if value is not None else '')
        doc.save(output_path)
    except Exception as e:
        raise ValueError(f"Erro ao substituir placeholders no documento Word: {e}")


def format_cpf(cpf):
    if cpf and cpf.isdigit():
        cpf = "{:0>11}".format(int(cpf))
        return re.sub(r"(\d{3})(\d{3})(\d{3})(\d{2})", r"\1.\2.\3-\4", cpf)
    return cpf


def format_cnpj(cnpj):
    if cnpj and cnpj.isdigit():
        cnpj = "{:0>14}".format(int(cnpj))
        return re.sub(r"(\d{2})(\d{3})(\d{3})(\d{4})(\d{2})", r"\1.\2.\3/\4-\5", cnpj)
    return cnpj


def clean_directory(directory):
    """Esvazia a pasta especificada."""
    for file_name in os.listdir(directory):
        file_path = os.path.join(directory, file_name)
        try:
            if os.path.isfile(file_path) or os.path.islink(file_path):
                os.unlink(file_path)
            elif os.path.isdir(file_path):
                os.rmdir(file_path)
        except Exception as e:
            print(f'Falha ao deletar {file_path}. Motivo: {e}')


def delete_zip_files(directory):
    """Deleta todos os arquivos ZIP no diretório especificado."""
    zip_files = find_zip_files(directory)
    for zip_file in zip_files:
        try:
            os.remove(zip_file)
            #print(f"Arquivo ZIP deletado: {zip_file}")
        except Exception as e:
            print(f"Falha ao deletar {zip_file}. Motivo: {e}")


def main():
    zip_dir = r'C:\Users\WinUser\Documents\TermoGarantiaSeggna\xml'
    extract_to = r'C:\Users\WinUser\Documents\TermoGarantiaSeggna\xml\extracted'
    doc_path = r'C:\Users\WinUser\Documents\TermoGarantiaSeggna\termoGarantia\TERMO DE GARANTIA.docx'

    os.makedirs(extract_to, exist_ok=True)

    zip_files = find_zip_files(zip_dir)
    if zip_files:
        for zip_path in zip_files:
            try:
                xml_paths = extract_xml_from_zip(zip_path, extract_to)
                if xml_paths:
                    for xml_path in xml_paths:
                        try:
                            data = parse_xml_data(xml_path)

                            if data:
                                # Definindo a data atual para o termo com mês em português
                                current_date = datetime.now()
                                replacements = {
                                    "_NOME_": data.get('Nome', '').upper(),
                                    "_NACIONALIDADE_": "Brasileiro",
                                    "_CPF/CNPJ_": data.get('Identificacao', ''),
                                    "_ENDERECO_": data.get('Endereco', ''),
                                    "_CEP_": data.get('CEP', ''),
                                    "_DIA_": current_date.strftime("%d"),
                                    "_MES_": current_date.strftime("%B").upper(),
                                    "_ANO_": current_date.strftime("%Y"),
                                    "_MODELOGPS_": data.get('ModeloGPS', ''),
                                    "_NUMSERIE_": data.get('NumeroSerie', '')
                                }

                                # Criando nome de arquivo baseado no nome do destinatário e número da nota fiscal
                                nome_sanitizado = data['Nome'].replace(' ', '_').replace('.', '').replace(',', '').upper()
                                numero_nota_fiscal = data['NumeroNotaFiscal']
                                doc_output_path = os.path.join(r'C:\Users\WinUser\Documents\TermoGarantiaSeggna',
                                                               f'{nome_sanitizado}_{numero_nota_fiscal}.docx')
                                pdf_output_path = os.path.join(r'C:\Users\WinUser\Documents\TermoGarantiaSeggna',
                                                               f'{nome_sanitizado}_{numero_nota_fiscal}.pdf')

                                replace_placeholders(doc_path, doc_output_path, replacements)
                                print(f"Documento Word atualizado salvo em: {doc_output_path}")

                                # Convertendo o documento Word para PDF
                                convert(doc_output_path, pdf_output_path)
                                print(f"Documento PDF salvo em: {pdf_output_path}")

                            else:
                                print(f"Falha ao extrair dados do XML do arquivo: {os.path.basename(xml_path)}")

                        except ValueError as ve:
                            print(f"Erro ao processar o arquivo {os.path.basename(xml_path)}: {ve}")
                        except Exception as e:
                            print(f"Erro inesperado ao processar o arquivo {os.path.basename(xml_path)}: {e}")

                else:
                    print(f"Nenhum arquivo XML encontrado no ZIP: {os.path.basename(zip_path)}")

            except ValueError as ve:
                print(f"Erro ao processar o arquivo {os.path.basename(zip_path)}: {ve}")
            except Exception as e:
                print(f"Erro inesperado ao processar o arquivo {os.path.basename(zip_path)}: {e}")

            # Esvaziando a pasta extraída para o próximo arquivo ZIP
            clean_directory(extract_to)

        # Deletar todos os arquivos ZIP após processar
        delete_zip_files(zip_dir)
    else:
        print("Nenhum arquivo ZIP encontrado no diretório.")


if __name__ == "__main__":
    main()
